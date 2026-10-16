import os, re, sys, collections
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx import Document
import builder
ROOT = builder.ROOT
BASE_LEN = len("C:\\Users\\Goutham\\Desktop\\carrier_new\\")

BAD = [r"\bNone\b", r"\[\]", r"\{", r"\}", r"if False", r"TODO", r"\?\?", r"nan\b", r"\bundefined\b", r"<\w+>", r"\s{3,}\S"]
problems = collections.defaultdict(list)
n = 0; total_paras = 0
titles = collections.Counter()
for dp, dn, fn in os.walk(ROOT):
    rel = os.path.relpath(dp, ROOT).replace(os.sep, "/")
    docs = [f for f in fn if f.endswith(".docx")]
    if "Abroad_Jobs" == os.path.basename(dp):
        if sorted(docs) != ["1_Top_10_Countries.docx", "2_Job_Search_Apps_by_Country.docx", "3_Steps_to_Go_Abroad.docx"]:
            problems["abroad_incomplete"].append(rel)
    for f in fn:
        p = os.path.join(dp, f)
        relp = os.path.relpath(p, ROOT)
        if BASE_LEN + len(relp) > 245:
            problems["path_too_long"].append((BASE_LEN + len(relp), relp))
        if not f.endswith(".docx"):
            problems["non_docx"].append(relp); continue
        n += 1
        try:
            d = Document(p)
        except Exception as e:
            problems["cannot_open"].append((relp, str(e))); continue
        paras = [x.text for x in d.paragraphs]
        total_paras += len(paras)
        title = paras[0] if paras else ""
        titles[(os.path.dirname(relp), title)] += 1
        if len([x for x in paras if x.strip()]) < 5:
            problems["too_short"].append(relp)
        heads = [x.text for x in d.paragraphs if x.style.name.startswith("Heading")]
        if len(heads) < 1:
            problems["no_headings"].append(relp)
        txt = "\n".join(paras)
        for t in d.tables:
            for r_i, row in enumerate(t.rows):
                for c in row.cells:
                    if not c.text.strip():
                        problems["empty_cell"].append(relp)
                    txt += "\n" + c.text
        for b in BAD:
            m = re.search(b, txt)
            if m:
                problems["bad_text:" + b].append((relp, txt[max(0, m.start()-40):m.end()+40].replace("\n", " | ")))
        # heading followed directly by heading (empty section)
        seq = [(x.style.name.startswith("Heading 1"), x.text) for x in d.paragraphs if x.text.strip()]
        for i in range(len(seq) - 1):
            if seq[i][0] and seq[i+1][0]:
                # allowed if next content is a table (tables not in paragraphs) -> check skipped
                pass
for k, v in titles.items():
    if v > 1:
        problems["duplicate_title_same_folder"].append(k)

# course folders without Abroad_Jobs
course_dirs_missing = []
for dp, dn, fn in os.walk(ROOT):
    rel = os.path.relpath(dp, ROOT).replace(os.sep, "/")
    if "Abroad_Jobs" in rel.split("/"):
        continue
    if any(f.endswith(".docx") for f in fn) and "Abroad_Jobs" not in dn:
        course_dirs_missing.append(rel)

print("DOCS:", n, " paragraphs:", total_paras)
for k, v in problems.items():
    print(f"\n### {k}: {len(v)}")
    for x in v[:12]:
        print("   ", x)
print("\n### Folders with docs but NO Abroad_Jobs:", len(course_dirs_missing))
for x in sorted(course_dirs_missing):
    print("   ", x)
