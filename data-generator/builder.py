"""Shared helpers to build consistent Word files for the career guide."""
import os
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.environ.get("CGT_OUT", os.path.join(os.path.dirname(os.path.abspath(__file__)), "output", "Career_Guidance_Telangana"))
NAVY = RGBColor(0x1F, 0x3A, 0x5F)
GREY = RGBColor(0x59, 0x59, 0x59)

NOTE = ("Note: Fees, seats, age limits, exam patterns and vacancies change every year. "
        "Always confirm with the latest official notification / website before applying. "
        "Information prepared as a guide (Oct 2026).")

COUNT = 0
RECORDS = []


def _shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def _base_doc(title, subtitle=None):
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(2)
        s.top_margin = s.bottom_margin = Cm(1.8)
    st = d.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)
    for lvl, size in ((1, 15), (2, 12.5)):
        h = d.styles[f"Heading {lvl}"]
        h.font.name = "Calibri"
        h.font.size = Pt(size)
        h.font.color.rgb = NAVY
        h.font.bold = True
    t = d.add_paragraph()
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(20)
    r.font.color.rgb = NAVY
    p = d.add_paragraph()
    r = p.add_run(subtitle or "Career Guidance – Telangana (India)")
    r.italic = True
    r.font.color.rgb = GREY
    # rule
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "8"); b.set(qn("w:space"), "4"); b.set(qn("w:color"), "1F3A5F")
    bdr.append(b); pPr.append(bdr)
    return d


def table(d, headers, rows):
    t = d.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        _shade(c, "1F3A5F")
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j, v in enumerate(row):
            cells[j].text = str(v)
            if i % 2:
                _shade(cells[j], "EEF3F8")
    d.add_paragraph()


def make(rel_path, title, sections, subtitle=None, intro=None):
    """sections: list of (heading, content) where content is
       list[str] -> bullets, str -> paragraph, ('table', headers, rows) -> table,
       or list mixing these."""
    global COUNT
    rec_secs = []
    for heading, content in sections:
        items = content if isinstance(content, list) else [content]
        blocks = []
        for it in items:
            if isinstance(it, tuple) and it[0] == "table":
                blocks.append({"t": "table", "h": [str(x) for x in it[1]], "r": [[str(c) for c in row] for row in it[2]]})
            elif isinstance(it, tuple) and it[0] == "h2":
                blocks.append({"t": "h2", "x": it[1]})
            elif isinstance(it, tuple) and it[0] == "p":
                blocks.append({"t": "p", "x": it[1]})
            elif isinstance(it, tuple) and it[0] == "b":
                blocks.append({"t": "li", "x": it[1]})
            elif isinstance(content, list):
                blocks.append({"t": "li", "x": it})
            else:
                blocks.append({"t": "p", "x": it})
        rec_secs.append({"h": heading, "b": blocks})
    RECORDS.append({"path": rel_path, "title": title, "sub": subtitle or "Career Guidance – Telangana (India)", "intro": intro, "s": rec_secs})
    d = _base_doc(title, subtitle)
    if intro:
        d.add_paragraph(intro)
    for heading, content in sections:
        d.add_heading(heading, level=1)
        items = content if isinstance(content, list) else [content]
        for it in items:
            if isinstance(it, tuple) and it[0] == "table":
                table(d, it[1], it[2])
            elif isinstance(it, tuple) and it[0] == "h2":
                d.add_heading(it[1], level=2)
            elif isinstance(it, tuple) and it[0] == "p":
                d.add_paragraph(it[1])
            elif isinstance(it, tuple) and it[0] == "b":
                d.add_paragraph(it[1], style="List Bullet")
            elif isinstance(content, list):
                d.add_paragraph(it, style="List Bullet")
            else:
                d.add_paragraph(it)
    p = d.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run(NOTE)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    full = os.path.join(ROOT, rel_path)
    os.makedirs(os.path.dirname(full), exist_ok=True)
    d.save(full)
    COUNT += 1


ANY_DEGREE_GOVT = ("Also eligible for ALL 'any degree' government exams – UPSC Civil Services, TGPSC Group I/II/III/IV, "
                   "SSC CGL, IBPS/SBI PO & Clerk, RBI Grade B/Assistant, RRB NTPC (graduate), TG Police SI, LIC AAO, CDS/AFCAT. "
                   "See folder 07_After_Degree_Govt_and_Private_Jobs.")


def career_doc(rel_path, title, info, subtitle=None):
    """Standard course document. info keys (all optional):
    about, duration, eligibility, admission, colleges, subjects, pass_rules,
    skills, private [(role, where, salary)], govt [(post, recruiter, eligibility)],
    self_emp, higher, route, tips, extra [(heading, content)], any_degree(bool)"""
    s = []
    if info.get("about"):
        s.append(("What is this course?", info["about"]))
    det = []
    if info.get("duration"):
        det.append(f"Duration: {info['duration']}")
    det += info.get("eligibility", [])
    if det:
        s.append(("Duration & eligibility", det))
    if info.get("admission"):
        s.append(("How to get admission (Telangana)", info["admission"]))
    if info.get("colleges"):
        s.append(("Where to study (Telangana examples)", info["colleges"]))
    if info.get("subjects"):
        s.append(("Main subjects", info["subjects"]))
    if info.get("pass_rules"):
        s.append(("Pass percentage / rules", info["pass_rules"]))
    if info.get("skills"):
        s.append(("Extra skills to learn alongside", info["skills"]))
    for h, c in info.get("extra", []):
        s.append((h, c))
    if info.get("private"):
        s.append(("PRIVATE JOBS after this course", [("table", ["Job role", "Where you work", "Starting salary (approx/month)"],
                                                       [list(r) for r in info["private"]])]))
    if info.get("govt") or info.get("any_degree"):
        c = []
        if info.get("govt"):
            c.append(("table", ["Govt post", "Recruiter", "Eligibility / note"], [list(r) for r in info["govt"]]))
        if info.get("any_degree"):
            c.append(("p", ANY_DEGREE_GOVT))
        s.append(("GOVERNMENT JOBS after this course", c))
    if info.get("self_emp"):
        s.append(("Self-employment / business options", info["self_emp"]))
    if info.get("higher"):
        s.append(("Higher studies after this", info["higher"]))
    if info.get("route"):
        s.append(("Student route map (step by step)", info["route"]))
    if info.get("tips"):
        s.append(("Tips", info["tips"]))
    make(rel_path, title, s, subtitle=subtitle)


def folder(rel_path):
    os.makedirs(os.path.join(ROOT, rel_path), exist_ok=True)
